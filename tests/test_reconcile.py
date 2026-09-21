"""Engine, ingestion, report and CLI tests. The example accounts are checked against
expected.json, which make_examples.py writes from how each case was BUILT, not from
the engine's own output."""

import json
import os
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from reconcile import __main__ as cli  # noqa: E402
from reconcile.config import ConfigError, get_account, load_accounts  # noqa: E402
from reconcile.engine import categorise_deduction, classify, reconcile  # noqa: E402
from reconcile.ingest import IngestError, ingest, load_table  # noqa: E402
from reconcile.report import write_report  # noqa: E402

EXAMPLES = ROOT / "examples"


@pytest.fixture(scope="session", autouse=True)
def make_examples():
    """Regenerate the example files so tests never depend on stale committed copies."""
    import subprocess
    subprocess.run([sys.executable, str(EXAMPLES / "make_examples.py")], check=True, capture_output=True)


def run_account(code):
    account = get_account(code)
    folder = EXAMPLES / code
    inv, inv_rej, inv_seen = ingest(folder / "invoices.csv", account.invoices, account.po_regex)
    pay, pay_rej, pay_seen = ingest(folder / "payments.csv", account.payments, account.po_regex,
                                    negate=account.payments_are_negative)
    result = reconcile(inv, pay, account)
    counts = {"invoice_seen": inv_seen, "invoice_used": len(inv), "invoice_rejected": len(inv_rej),
              "payment_seen": pay_seen, "payment_used": len(pay), "payment_rejected": len(pay_rej)}
    return account, result, inv_rej, pay_rej, counts


# ------------------------------------------------------------------ classification
@pytest.mark.parametrize("invoice,net,lines,tol,expected", [
    (50000, 50000, 1, 1, "Paid"),
    (50000, 49999, 1, 1, "Paid"),            # within one cent
    (50000, 49998, 1, 1, "Short Pay"),       # two cents short
    (50000, 49995, 1, 5, "Paid"),            # a five-cent tolerance
    (50000, 0, 0, 1, "Unpaid"),
    (50000, 60000, 1, 1, "Overpaid"),
    (50000, 100000, 2, 1, "Duplicate"),
    (50000, 100000, 1, 1, "Overpaid"),       # one big payment is an overpayment, not a duplicate
    (50000, 55000, 2, 1, "Overpaid"),        # two payments, but not a clean multiple
])
def test_classify(invoice, net, lines, tol, expected):
    assert classify(invoice, net, lines, tol) == expected


def test_deduction_categories_use_first_matching_rule():
    rules = {"Freight": "freight|shipping", "Fees": "fee|charge"}
    assert categorise_deduction("Freight deduction", rules) == "Freight"
    assert categorise_deduction("Processing FEE", rules) == "Fees"
    assert categorise_deduction("something else", rules) == "Uncategorised"
    assert categorise_deduction("", rules) == "Uncategorised"


# ------------------------------------------------------------------ example accounts
@pytest.mark.parametrize("code", ["AC1000", "AC2000"])
def test_example_account_matches_the_answer_it_was_built_with(code):
    expected = json.loads((EXAMPLES / code / "expected.json").read_text())
    account, result, inv_rej, pay_rej, counts = run_account(code)
    got = {s: result.summary["status_counts"].get(s, 0) for s in expected["status_counts"]}
    assert got == expected["status_counts"]
    assert int((result.invoices["Explained by deductions"] == "Yes").sum()) == expected["explained_by_deductions"]
    assert result.summary["unapplied_count"] == expected["unapplied_lines"]
    assert result.summary["deduction_count"] == expected["deduction_lines"]
    assert len(inv_rej) == expected["rejected_invoice_rows"]
    assert len(pay_rej) == expected["rejected_payment_rows"]


@pytest.mark.parametrize("code", ["AC1000", "AC2000"])
def test_every_input_row_is_accounted_for(code):
    expected = json.loads((EXAMPLES / code / "expected.json").read_text())
    _, _, _, _, c = run_account(code)
    assert c["invoice_seen"] == c["invoice_used"] + c["invoice_rejected"] == expected["invoice_rows"]
    assert c["payment_seen"] == c["payment_used"] + c["payment_rejected"] == expected["payment_rows"]


def test_rejected_rows_say_why_and_where():
    _, _, inv_rej, pay_rej, _ = run_account("AC1000")
    reasons = " | ".join(pay_rej["reason"])
    assert "missing PO" in reasons and "not a number" in reasons
    assert "TBD" in inv_rej.iloc[0]["reason"]
    assert (pay_rej["source_row"] > 1).all()          # points at the row in the original file


def test_ac2000_header_is_on_row_4_and_payments_are_negated():
    account = get_account("AC2000")
    raw = load_table(EXAMPLES / "AC2000" / "payments.csv", account.payments.header_row)
    assert "Invoice Ref" in raw.columns
    pay, _, _ = ingest(EXAMPLES / "AC2000" / "payments.csv", account.payments, "", negate=True)
    assert (pay["amount_cents"] > 0).sum() > (pay["amount_cents"] < 0).sum()    # mostly money received


def test_repairs_excel_damaged_pos_so_they_still_match():
    _, result, _, _, _ = run_account("AC2000")
    # 3 invoices were unpaid by design; nothing may be "unpaid" just because its PO was 2.088E+9
    assert result.summary["status_counts"]["Unpaid"] == 3


def test_money_is_exact_in_cents():
    account = get_account("AC1000")
    inv = pd.DataFrame({"source_row": [2], "po_raw": ["A"], "po_key": ["A"], "amount_cents": [10],
                        "date": [""], "description": [""]})
    pay = pd.DataFrame({"source_row": [2, 3, 4], "po_raw": ["A"] * 3, "po_key": ["A"] * 3,
                        "amount_cents": [1, 1, 8], "date": [""] * 3, "description": [""] * 3})
    res = reconcile(inv, pay, account)
    assert res.invoices.iloc[0]["Net paid"] == 0.10 and res.invoices.iloc[0]["Status"] == "Paid"


# ------------------------------------------------------------------ errors and config
def test_unknown_account_lists_the_known_ones():
    with pytest.raises(ConfigError, match="AC1000"):
        get_account("NOPE")


def test_missing_column_error_names_the_columns_that_exist(tmp_path):
    f = tmp_path / "inv.csv"
    f.write_text("Order,Amt\n1,5\n")
    spec = get_account("AC1000").invoices
    with pytest.raises(IngestError, match="Columns in the file: Order, Amt"):
        ingest(f, spec)


def test_unsupported_file_type_is_refused(tmp_path):
    f = tmp_path / "x.pdf"
    f.write_text("x")
    with pytest.raises(IngestError, match="Unsupported"):
        load_table(f)


def test_config_rejects_a_negative_tolerance(tmp_path):
    cfg = tmp_path / "a.toml"
    cfg.write_text('[accounts.X]\n[accounts.X.invoices]\npo="a"\namount="b"\n'
                   '[accounts.X.payments]\npo="a"\namount="b"\n[accounts.X.rules]\ntolerance=-1\n')
    with pytest.raises(ConfigError, match="negative"):
        load_accounts(cfg)


# ------------------------------------------------------------------ report + CLI
def test_excel_report_has_the_sheets_and_confirms_row_accounting(tmp_path):
    account, result, inv_rej, pay_rej, counts = run_account("AC1000")
    out = write_report(tmp_path / "r.xlsx", account, result, inv_rej, pay_rej, counts)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Invoices", "Unapplied payments", "Deductions", "Rejected rows"]
    assert wb["Summary"]["A26"].value == "Every row accounted for"
    assert wb["Summary"]["B5"].value.startswith("=COUNTIF")         # summary is live formulas
    assert wb["Invoices"].max_row == len(result.invoices) + 1


def test_cli_runs_and_writes_a_report(tmp_path, capsys):
    out = tmp_path / "AC2000.xlsx"
    code = cli.main(["--account", "AC2000", "--invoices", str(EXAMPLES / "AC2000" / "invoices.csv"),
                     "--payments", str(EXAMPLES / "AC2000" / "payments.csv"), "--out", str(out)])
    assert code == 0 and out.exists()
    assert "Duplicate" in capsys.readouterr().out


def test_cli_reports_a_bad_column_with_exit_code_2(tmp_path, capsys):
    bad = tmp_path / "payments.csv"
    bad.write_text("Ref,Value\nX,1\n")
    code = cli.main(["--account", "AC1000", "--invoices", str(EXAMPLES / "AC1000" / "invoices.csv"),
                     "--payments", str(bad), "--out", str(tmp_path / "x.xlsx")])
    assert code == 2 and "not found" in capsys.readouterr().err


def test_cli_lists_accounts(capsys):
    assert cli.main(["--list-accounts"]) == 0
    assert "AC1000" in capsys.readouterr().out
