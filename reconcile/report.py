"""Write the reconciliation result to an Excel workbook."""

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .normalize import dollars

NAVY = "1F3864"
STATUS_FILL = {"Unpaid": "F8CBAD", "Short Pay": "FFE699", "Overpaid": "BDD7EE",
               "Duplicate": "D9B3FF", "Paid": "C6E0B4"}
STATUSES = ["Paid", "Unpaid", "Short Pay", "Overpaid", "Duplicate"]
MONEY = '$#,##0.00;($#,##0.00);-'


def _style_sheet(ws, money_cols=(), widths=None):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
    for name in money_cols:
        idx = next((c.column for c in ws[1] if c.value == name), None)
        if idx:
            for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                row[0].number_format = MONEY
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = max((len(str(c.value)) for c in col[1:] if c.value is not None), default=8)
        header = len(str(col[0].value or ""))
        # leave room for the filter arrow on the header cell
        ws.column_dimensions[letter].width = (widths or {}).get(
            letter, min(34, max(header + 6, longest + 3, 11)))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_report(path, account, result, invoice_rejects, payment_rejects, counts):
    """counts: dict with the row-accounting numbers from ingestion."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    inv = result.invoices
    n = len(inv)

    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        pd.DataFrame().to_excel(xl, sheet_name="Summary")           # placeholder, filled below
        inv.to_excel(xl, sheet_name="Invoices", index=False)
        un = result.unapplied.assign(**{"Amount": result.unapplied["amount_cents"].map(dollars)})
        un[["source_row", "po_raw", "po_key", "Amount", "date", "description"]].rename(columns={
            "source_row": "File row", "po_raw": "Original PO", "po_key": "PO", "date": "Date",
            "description": "Description"}).to_excel(xl, sheet_name="Unapplied payments", index=False)
        ded = result.deductions.assign(Amount=result.deductions["amount_cents"].map(lambda c: dollars(-c)))
        ded[["source_row", "po_key", "category", "Amount", "description", "matched_invoice"]].rename(columns={
            "source_row": "File row", "po_key": "PO", "category": "Category",
            "description": "Description", "matched_invoice": "Has invoice"}).to_excel(
            xl, sheet_name="Deductions", index=False)
        pd.concat([invoice_rejects.assign(file="invoices"), payment_rejects.assign(file="payments")],
                  ignore_index=True).rename(columns={
            "source_row": "File row", "po_raw": "PO", "amount_raw": "Amount", "reason": "Reason",
            "file": "File"}).to_excel(xl, sheet_name="Rejected rows", index=False)
        wb = xl.book

        _style_sheet(wb["Invoices"], money_cols=["Invoice amount", "Received", "Deductions",
                                                 "Net paid", "Difference"])
        for row in wb["Invoices"].iter_rows(min_row=2, min_col=10, max_col=10):
            fill = STATUS_FILL.get(row[0].value)
            if fill:
                row[0].fill = PatternFill("solid", fgColor=fill)
        _style_sheet(wb["Unapplied payments"], money_cols=["Amount"])
        _style_sheet(wb["Deductions"], money_cols=["Amount"])
        _style_sheet(wb["Rejected rows"])

        ws = wb["Summary"]
        ws.delete_rows(1, ws.max_row)
        ws.sheet_view.showGridLines = False
        ws["A1"] = f"Reconciliation: {account.code}, {account.name}"
        ws["A1"].font = Font(bold=True, size=15, color=NAVY, name="Arial")
        ws["A2"] = (f"Tolerance ${account.tolerance_cents / 100:.2f} | "
                    f"PO suffix rule: {account.po_regex or 'none'} | "
                    f"payments exported as negative: {'yes' if account.payments_are_negative else 'no'}")
        ws["A2"].font = Font(italic=True, size=9, color="595959", name="Arial")
        heads = ["Status", "Invoices", "Invoice amount", "Net paid", "Difference"]
        for c, h in enumerate(heads, start=1):
            cell = ws.cell(row=4, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(horizontal="center")
        last = n + 1
        for i, st in enumerate(STATUSES, start=5):
            ws.cell(row=i, column=1, value=st)
            ws.cell(row=i, column=2, value=f'=COUNTIF(Invoices!$J$2:$J${last},A{i})')
            ws.cell(row=i, column=3, value=f'=SUMIF(Invoices!$J$2:$J${last},A{i},Invoices!$D$2:$D${last})')
            ws.cell(row=i, column=4, value=f'=SUMIF(Invoices!$J$2:$J${last},A{i},Invoices!$G$2:$G${last})')
            ws.cell(row=i, column=5, value=f'=SUMIF(Invoices!$J$2:$J${last},A{i},Invoices!$H$2:$H${last})')
        ws.cell(row=10, column=1, value="Total").font = Font(bold=True, name="Arial", size=10)
        for c in range(2, 6):
            L = get_column_letter(c)
            ws.cell(row=10, column=c, value=f"=SUM({L}5:{L}9)").font = Font(bold=True, name="Arial", size=10)
        for r in range(5, 11):
            ws.cell(row=r, column=2).number_format = "#,##0"
            for c in (3, 4, 5):
                ws.cell(row=r, column=c).number_format = MONEY
            ws.cell(row=r, column=1).font = Font(name="Arial", size=10, bold=(r == 10))

        s = result.summary
        ws["A12"] = "Findings"
        ws["A12"].font = Font(bold=True, size=11, color=NAVY, name="Arial")
        findings = [
            ("Owed to you (Unpaid + Short Pay)", s["owed"], MONEY),
            ("Excess received (Overpaid + Duplicate)", s["excess"], MONEY),
            ("Unapplied payments (no invoice)", f'{s["unapplied_count"]} lines, ${s["unapplied_amount"]:,.2f}', None),
            ("Deductions taken", f'{s["deduction_count"]} lines, ${s["deduction_amount"]:,.2f}', None),
            ("Short payments explained by deductions",
             f'=COUNTIF(Invoices!$K$2:$K${last},"Yes")', "#,##0"),
        ]
        for i, (label, val, fmt) in enumerate(findings, start=13):
            ws.cell(row=i, column=1, value=label).font = Font(name="Arial", size=10)
            cell = ws.cell(row=i, column=2, value=val)
            cell.font = Font(name="Arial", size=10, bold=True)
            if fmt:
                cell.number_format = fmt

        ws["A19"] = "Row accounting"
        ws["A19"].font = Font(bold=True, size=11, color=NAVY, name="Arial")
        acct = [
            ("Invoice rows read", counts["invoice_seen"]),
            ("  used", counts["invoice_used"]),
            ("  rejected", counts["invoice_rejected"]),
            ("Payment rows read", counts["payment_seen"]),
            ("  used", counts["payment_used"]),
            ("  rejected", counts["payment_rejected"]),
        ]
        for i, (label, val) in enumerate(acct, start=20):
            ws.cell(row=i, column=1, value=label).font = Font(name="Arial", size=10)
            ws.cell(row=i, column=2, value=val).font = Font(name="Arial", size=10)
        ok = (counts["invoice_seen"] == counts["invoice_used"] + counts["invoice_rejected"]
              and counts["payment_seen"] == counts["payment_used"] + counts["payment_rejected"])
        ws["A26"] = "Every row accounted for" if ok else "ROW COUNT MISMATCH: check the input files"
        ws["A26"].font = Font(bold=True, name="Arial", size=10, color="2E7D32" if ok else "C62828")
        for col, w in zip("ABCDE", [44, 22, 18, 18, 18]):
            ws.column_dimensions[col].width = w
        wb.move_sheet("Summary", offset=-(len(wb.sheetnames) - 1))
        wb.active = 0
    return path
